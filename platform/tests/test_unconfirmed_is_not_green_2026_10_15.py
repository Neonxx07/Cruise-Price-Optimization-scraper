"""An UNCONFIRMED candidate must not look like a confirmed saving.

Found 2026-10-15 in the pre-run check. Since §1.6 (BUGS_AND_FIXES) a
self-declared UNCONFIRMED GoCCL candidate is kept out of the savings total,
but its row kept the OPTIMIZATION green in the GUI and in the Excel export.
After the 2026-10-12 review work, a candidate Carnival's review could not
confirm (REVIEW_NOT_CAPTURED, an invoice that does not add up - the TFPA
FunPlay line on DEMO20/DEMO16) is exactly such a row. Now pale yellow:
verify by hand.
"""

from openpyxl import load_workbook

from core.calculator import is_unconfirmed_candidate
from core.models import BookingResult, BookingStatus, CruiseLine
from services.excel_export import export_results_excel


def _r(note):
    return BookingResult(cruise_line=CruiseLine.GOCCL, booking_id="DEMO20",
                         status=BookingStatus.OPTIMIZATION, net_saving=66.0,
                         old_total=2162.0, new_total=2096.0, price_drop=66.0, note=note)


def test_the_unconfirmed_note_is_recognised():
    assert is_unconfirmed_candidate(_r("UNCONFIRMED candidate FZD [REVIEW: not confirmed - x]"))
    assert not is_unconfirmed_candidate(_r("saving 48.00 - confirmed on review"))


def test_the_gui_has_a_tint_of_its_own_that_is_not_green():
    from gui.windows import CruiseLinePanel

    tints = CruiseLinePanel._STATUS_TINTS
    assert tints["UNCONFIRMED"] != tints["OPTIMIZATION"]


def test_the_excel_row_is_yellow_not_green(tmp_path):
    path = tmp_path / "x.xlsx"
    export_results_excel([_r("UNCONFIRMED candidate"), _r("saving 48.00 confirmed")], path)
    ws = load_workbook(path).active
    fills = {ws.cell(row=i, column=13).value: ws.cell(row=i, column=1).fill.fgColor.rgb
             for i in (2, 3)}
    assert fills["UNCONFIRMED candidate"].endswith("FFFF99")
    assert fills["saving 48.00 confirmed"].endswith("C6EFCE")
